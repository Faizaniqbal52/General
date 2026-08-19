"""Excel and CSV export of the application table."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from .rows import COLUMNS

BAND_FILL = {"A+": "FFD8F0DC", "A": "FFE8F5E9", "B": "FFFFF8E1", "C": "FFFDF1E7", "SKIP": "FFF1F1F1"}
WIDTHS = {"priority": 9, "company": 24, "role": 34, "overall": 10, "technical": 11,
          "eligibility": 12, "domain": 12, "salary": 18, "location": 22, "mode": 12,
          "contact_name": 20, "contact_title": 20, "contact_email": 28,
          "contact_email_status": 16, "contact_link": 28, "resume": 40,
          "email_subject": 36, "email_body": 70, "linkedin_note": 46, "why": 40,
          "gaps": 40, "verdict": 60, "job_url": 34, "status": 16, "new": 6,
          "first_seen": 12, "source": 14, "searches": 50}


def write_csv(rows: list[dict[str, Any]], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow([label for _, label in COLUMNS])
        for row in rows:
            writer.writerow([row.get(key, "") for key, _ in COLUMNS])
    return path


def write_sheet(rows: list[dict[str, Any]], path: str | Path) -> Path:
    """Write .xlsx when openpyxl is available, otherwise fall back to .csv."""
    path = Path(path)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError:
        return write_csv(rows, path.with_suffix(".csv"))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Applications"

    header_font = Font(bold=True, color="FFFFFFFF")
    header_fill = PatternFill("solid", fgColor="FF263238")
    for index, (_, label) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", horizontal="left")
    sheet.freeze_panes = "C2"

    for r, row in enumerate(rows, start=2):
        fill = PatternFill("solid", fgColor=BAND_FILL.get(str(row.get("priority")), "FFFFFFFF"))
        for c, (key, _) in enumerate(COLUMNS, start=1):
            value = row.get(key, "")
            cell = sheet.cell(row=r, column=c, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=key in
                                       {"email_body", "verdict", "why", "gaps", "searches",
                                        "linkedin_note", "email_subject"})
            if key in {"priority", "overall", "technical", "eligibility", "domain"}:
                cell.fill = fill
            if key in {"job_url", "contact_link"} and isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.style = "Hyperlink"
            if key == "resume" and value:
                cell.hyperlink = f"file://{Path(value).resolve()}"
                cell.style = "Hyperlink"

    for index, (key, _) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS.get(key, 18)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 2)}"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
